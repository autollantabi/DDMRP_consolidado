#ExportarExcel
# Exporta core.ddmrp a Excel con el formato del "Manual Formato de Reporte de DDMRP".
# Las columnas se buscan por NOMBRE (no por letra), así el formato no se daña si cambia el orden.
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from config import get_engine_postgres_dwh

# --- 1. CONSULTA SQL ---
# 0.1 solo los activos para compra y venta (ACTIVO = 'SI' y ARTICULO_COMPRA = 'Y')
query = """
select * from core.ddmrp where "GRUPO" in ('LLANTAS', 'LLANTAS MOTO') AND "ACTIVO" = 'SI' AND "ARTICULO_COMPRA" = 'Y';
"""

# --- 2. NOMBRE DEL ARCHIVO ---
# Reporte de Pedidos (Marcas) (Empresa) (fecha de creación DD-MM-AAAA).xlsx
MARCAS  = "NEUMATICOS"
EMPRESA = "TODAS"
HOY     = date.today().strftime("%d-%m-%Y")
CARPETA = Path(r"C:\Users\plituma\Documents\REPORTES DDMRP") / HOY
output_file = CARPETA / f"Reporte de Pedidos ({MARCAS}) ({EMPRESA}) ({HOY}).xlsx"

# --- 3. COLUMNAS (nombres en core.ddmrp) ---
COL_ACTIVO        = "ACTIVO"
COL_ART_COMPRA    = "ARTICULO_COMPRA"
COL_MES_INV_TOTAL = "MES INV TOTAL SEGUN DEMANDA 90D"
COL_DIAS_INV_90D  = "DIAS INV 90D"
COL_DIAS_INV_1A   = "DIAS INV 1 AÑO"
COL_DIF_PICO      = "DIF. MAX AÑO vs TOT. AÑO"
COL_BARRAS        = "BARRAS"

# --- 4. COLORES ---
NARANJA       = PatternFill("solid", fgColor="FFC000")
ROJO          = PatternFill("solid", fgColor="FF0000")
VERDE         = PatternFill("solid", fgColor="92D050")
GRIS_CLARO    = PatternFill("solid", fgColor="D9D9D9")
GRIS_OSCURO   = PatternFill("solid", fgColor="A6A6A6")
AZUL          = PatternFill("solid", fgColor="9BC2E6")
AMARILLO      = PatternFill("solid", fgColor="FFFF00")
TOMATE_PALIDO = PatternFill("solid", fgColor="F8CBAD")

# 1. NARANJA: DIAS INV 90D < 68
# 2. ROJO:    DIAS INV 90D < 68  y  DIAS INV 1 AÑO < 275   (se pinta encima del naranja)
# 3. VERDE:   DIF. MAX AÑO vs TOT. AÑO = 'SI'              (se pinta al final, encima de los otros)
LIMITE_DIAS_90D = 68
LIMITE_DIAS_1A  = 275

# 4. columnas que se ocultan (0.1: ACTIVO y ARTICULO_COMPRA; 4.11: DIF. MAX AÑO vs TOT. AÑO)
COLUMNAS_OCULTAS = [COL_ACTIVO, COL_ART_COMPRA, COL_DIF_PICO]

# 5. color de columnas completas (se pinta después de las filas, como en el manual)
COLOR_COLUMNAS = {
    "DEMANDA MES - AÑO":    GRIS_CLARO,
    "DEMANDA MES 90D":      GRIS_CLARO,
    "EN STOCK":             GRIS_OSCURO,
    COL_MES_INV_TOTAL:      AZUL,
    "STOCK TOTAL":          AZUL,
    "NUEVO TAMAÑO PEDIDO":  AMARILLO,
    "PEDIDO SUGERIDO JP":   TOMATE_PALIDO,     # no existe en core.ddmrp: se salta si no está
}

# 7. columnas que se agrupan: cada línea agrupa desde su primera hasta su última columna
#    (las del medio son solo de referencia; las columnas deben estar seguidas en el reporte)
GRUPOS = [
    ("CODIGO_ITEM", "CODIGO PROVEEDOR", "COD. PROV1", "PROV1", "COD. PROV2", "PROV2"),
    ("LTF", "VF"),
    ("ADU SIN PICOS", "TAMAÑO BUFFER (TOG)"),
    ("NFP JP", "VALIDADOR"),                     # se salta si no están en core.ddmrp
]


def dar_formato(ws, df):
    """Aplica al Excel los pasos del manual: filas, columnas ocultas, colores, grupos."""
    columnas = list(df.columns)
    letra = {c: get_column_letter(i + 1) for i, c in enumerate(columnas)}
    n_cols = len(columnas)
    faltan = []

    # 1-3. pintar filas (la fila 1 es el encabezado: los datos empiezan en la 2)
    dias_90d = pd.to_numeric(df[COL_DIAS_INV_90D], errors="coerce")
    dias_1a = pd.to_numeric(df[COL_DIAS_INV_1A], errors="coerce")
    naranja = dias_90d < LIMITE_DIAS_90D
    rojo = naranja & (dias_1a < LIMITE_DIAS_1A)
    verde = df[COL_DIF_PICO] == "SI"
    for i in range(len(df)):
        if verde.iat[i]:
            relleno = VERDE
        elif rojo.iat[i]:
            relleno = ROJO
        elif naranja.iat[i]:
            relleno = NARANJA
        else:
            continue
        for celda in ws[i + 2][:n_cols]:
            celda.fill = relleno
    print(f"Filas pintadas: naranja {int((naranja & ~rojo & ~verde).sum())} | "
          f"rojo {int((rojo & ~verde).sum())} | verde {int(verde.sum())}")

    # 4. ocultar columnas
    for c in COLUMNAS_OCULTAS:
        if c in letra:
            ws.column_dimensions[letra[c]].hidden = True
        else:
            faltan.append(c)

    # 5. colorear columnas completas (encabezado y datos)
    for c, relleno in COLOR_COLUMNAS.items():
        if c not in letra:
            faltan.append(c)
            continue
        for (celda,) in ws.iter_rows(min_col=columnas.index(c) + 1, max_col=columnas.index(c) + 1,
                                     min_row=1, max_row=len(df) + 1):
            celda.fill = relleno

    # 7. agrupar columnas
    # cada grupo va desde su primera columna hasta la última (puede tener 2 o más nombres)
    for grupo in GRUPOS:
        desde, hasta = grupo[0], grupo[-1]
        if desde in letra and hasta in letra:
            ws.column_dimensions.group(letra[desde], letra[hasta], outline_level=1, hidden=False)
        else:
            faltan.extend(c for c in (desde, hasta) if c not in letra)

    if faltan:
        print("Columnas del manual que no están en el reporte (se saltan):", ", ".join(faltan))


# Hoja con la explicación de los colores de las filas (texto en A2:A4, cada uno con su color)
HOJA_EXPLICACION = "EXPLICACIÓN COLORES"
EXPLICACION_COLORES = [
    ("NARANJA = DIAS INV 90D MENOR A 68 DIAS", NARANJA),
    ("ROJO = DIAS INV 1 AÑO MENOR A 275 DIAS", ROJO),
    ("VERDE = DIF PICO VTAS ULT AÑO VS # VTAS ULT AÑO = SI", VERDE),
]


def agregar_explicacion(wb):
    """Agrega la hoja EXPLICACIÓN COLORES."""
    ws = wb.create_sheet(HOJA_EXPLICACION)
    for fila, (texto, relleno) in enumerate(EXPLICACION_COLORES, start=2):
        celda = ws.cell(row=fila, column=1, value=texto)
        celda.fill = relleno
    ws.column_dimensions["A"].width = 75


try:
    print("Conectando a PostgreSQL y ejecutando consulta...")
    engine = get_engine_postgres_dwh()

    # Cargar el resultado directamente en un DataFrame de Pandas
    df = pd.read_sql_query(query, engine)

    # 0.2 MES INV TOTAL: 9999 -> INFINITO
    df[COL_MES_INV_TOTAL] = df[COL_MES_INV_TOTAL].astype(object).where(
        pd.to_numeric(df[COL_MES_INV_TOTAL], errors="coerce") != 9999, "INFINITO")

    # 8. ordenar por BARRAS de la A a la Z (vacíos al final)
    df = df.sort_values(COL_BARRAS, na_position="last", kind="stable").reset_index(drop=True)

    print(f"Exportando {len(df)} filas a Excel...")
    CARPETA.mkdir(parents=True, exist_ok=True)

    # Exportar a Excel (index=False evita exportar el número de fila implícito de Pandas)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)
        dar_formato(writer.sheets["Sheet1"], df)
        agregar_explicacion(writer.book)

    print(f"¡Éxito! Archivo guardado correctamente como '{output_file}'")

except Exception as e:
    print(f"Ocurrió un error durante el proceso: {e}")
